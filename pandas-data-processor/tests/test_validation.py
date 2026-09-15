"""Validation must fail visibly and prevent report publication."""
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import ROOT
from errors import SkillError
from table_io import read_table, write_tables
from validation import compare_frames, validate_frame, validate_result_frame


def test_valid_schema_and_result(fixtures, schema, contract):
    raw = read_table(fixtures / "sales.csv")
    cleaned, errors = validate_frame(raw, schema)
    assert len(errors) == 0
    assert cleaned.employee_id.tolist() == ["001234", "001235"]
    assert cleaned.amount.sum() == 30
    assert validate_result_frame(cleaned, contract)
    assert raw.amount.tolist() == ["10.25", "19.75"]


def test_missing_required_column(schema):
    with pytest.raises(SkillError) as caught:
        validate_frame(pd.DataFrame({"employee_id": ["001234"]}), schema)
    assert caught.value.code == "SCHEMA_MISMATCH"


def test_invalid_values_retained(fixtures, schema):
    raw = read_table(fixtures / "invalid.csv")
    cleaned, errors = validate_frame(raw, schema)
    assert len(cleaned) == 0
    assert {"DATE_PARSE_ERROR", "NUMERIC_PARSE_ERROR", "DUPLICATE_KEY"}.issubset(set(errors.error_code))
    numeric = errors.loc[errors.error_code == "NUMERIC_PARSE_ERROR"].iloc[0]
    assert numeric.original_value == "oops"
    assert numeric.source_row == 3
    assert numeric.source_file == str((fixtures / "invalid.csv").resolve())


@pytest.mark.parametrize("rule,value", [
    ({"dtype": "number", "min": 0}, "-1"),
    ({"dtype": "number", "max": 10}, "11"),
    ({"dtype": "integer"}, "1.5"),
    ({"dtype": "boolean"}, "yes"),
    ({"dtype": "string", "enum": ["A"]}, "B"),
    ({"dtype": "number"}, "inf"),
])
def test_constraints(rule, value):
    valid, errors = validate_frame(pd.DataFrame({"a": [value]}), {"required_columns": {"a": rule}})
    assert valid.empty and len(errors) >= 1


def test_missing_not_filled_and_optional_checked():
    valid, errors = validate_frame(pd.DataFrame({"a": [pd.NA], "b": ["wrong"]}),
        {"required_columns": {"a": "number"}, "optional_columns": {"b": "integer"}})
    assert len(errors) == 1 and valid.empty
    valid, errors = validate_frame(pd.DataFrame({"a": [pd.NA]}), {"required_columns": {"a": "number"}})
    assert len(valid) == 1 and pd.isna(valid.iloc[0]["a"])


def test_total_mismatch_and_tolerance():
    left = pd.DataFrame({"amount": ["0.1", "0.2"]})
    exact = pd.DataFrame({"amount": ["0.3"]})
    assert compare_frames(left, exact, ["amount"])["amount"]["pass"]
    different = pd.DataFrame({"amount": ["0.31"]})
    with pytest.raises(SkillError) as caught:
        compare_frames(left, different, ["amount"])
    assert caught.value.code == "TOTAL_MISMATCH"
    assert compare_frames(left, different, ["amount"], atol="0.01")


def test_precision_loss_is_an_exception():
    value = "0.12345678901234567890123456789"
    valid, errors = validate_frame(pd.DataFrame({"amount": [value]}), {"required_columns": {"amount": "number"}})
    assert valid.empty
    assert errors.iloc[0].original_value == value


def test_long_decimal_total_is_exact():
    from validation import total
    frame = pd.DataFrame({"amount": ["123456789012345678901234567890.01", "0.02"]})
    assert str(total(frame, "amount")) == "123456789012345678901234567890.03"


@pytest.mark.parametrize("values", [[None], ["bad"], ["inf"]])
def test_total_never_skips_bad_amounts(values):
    with pytest.raises(SkillError):
        compare_frames(pd.DataFrame({"amount": values}), pd.DataFrame({"amount": ["0"]}), ["amount"])


@pytest.mark.parametrize("contract,code", [
    ({"row_count": 1}, "ROW_COUNT_MISMATCH"),
    ({"null_counts": {"amount": 1}}, "VALIDATION_FAILED"),
    ({"unique_ids": {"employee_id": ["001234"]}}, "VALIDATION_FAILED"),
    ({"totals": {"amount": 31}}, "TOTAL_MISMATCH"),
    ({"exception_count": 1}, "VALIDATION_FAILED"),
    ({"required_columns": ["missing"]}, "SCHEMA_MISMATCH"),
])
def test_result_contract_failures(fixtures, contract, code):
    with pytest.raises(SkillError) as caught:
        validate_result_frame(read_table(fixtures / "sales.csv"), contract)
    assert caught.value.code == code


def test_csv_and_xlsx_workflows(cli, fixtures, tmp_path):
    schema = ROOT / "specs" / "schema.example.json"
    contract = fixtures / "result.json"
    cli("validate_schema", fixtures / "sales.csv", "--schema", schema)
    for suffix in ("csv", "xlsx"):
        output = tmp_path / f"report.{suffix}"
        args = ["--group-by", "department", "--sum-columns", "amount"] if suffix == "xlsx" else []
        report = cli("export_report", fixtures / "sales.csv", "--schema", schema,
                     "--contract", contract, "--output", output, *args)
        assert report["valid_rows"] == 2 and report["exception_count"] == 0
        sheet = ["--sheet", "Cleaned"] if suffix == "xlsx" else []
        cli("validate_result", output, "--contract", contract, *sheet)
        after_sheet = ["--after-sheet", "Summary"] if suffix == "xlsx" else []
        cli("compare_totals", fixtures / "sales.csv", output, "--columns", "amount", *after_sheet)
        if suffix == "xlsx":
            assert load_workbook(output).sheetnames == ["Raw", "Cleaned", "Summary", "Exceptions"]


def test_failed_validation_writes_no_report(cli, fixtures, tmp_path):
    output = tmp_path / "blocked.xlsx"
    result = cli("export_report", fixtures / "invalid.csv", "--schema", ROOT / "specs" / "schema.example.json",
                 "--contract", fixtures / "result.json", "--output", output, expected=2)
    assert result["status"] == "error"
    assert not output.exists()
    contract = tmp_path / "wrong.json"
    contract.write_text(json.dumps({"totals": {"amount": 999}}), encoding="utf-8")
    result = cli("export_report", fixtures / "sales.csv", "--schema", ROOT / "specs" / "schema.example.json",
                 "--contract", contract, "--output", output, expected=2)
    assert result["error"]["code"] == "TOTAL_MISMATCH" and not output.exists()


def test_export_safety_and_literal_formulas(tmp_path):
    source = tmp_path / "source.csv"
    source.write_text("id\n001\n", encoding="utf-8")
    frame = pd.DataFrame({"id": ["001"], "text": ["=1+1"]})
    with pytest.raises(SkillError):
        write_tables({"Data": frame}, source, sources=[source], overwrite=True)
    assert source.read_text(encoding="utf-8") == "id\n001\n"
    output = tmp_path / "safe.xlsx"
    write_tables({"Data": frame}, output)
    book = load_workbook(output, data_only=False)
    assert book["Data"]["A2"].value == "001"
    assert book["Data"]["B2"].data_type == "s"
    assert book["Data"]["B2"].value == "=1+1"
    with pytest.raises(SkillError):
        write_tables({"Data": frame}, output)


def test_safe_csv_is_opt_in_and_counted(tmp_path):
    frame = pd.DataFrame({"text": ["=1+1", "normal"]})
    path = tmp_path / "safe.csv"
    result = write_tables({"Data": frame}, path, spreadsheet_safe_csv=True)
    assert result["csv_formula_cells_escaped"] == 1
    assert read_table(path).text.tolist() == ["'=1+1", "normal"]


@pytest.mark.parametrize("contract", [
    {"row_count": True}, {"duplicate_count": -1}, {"required_columns": "id"},
    {"null_counts": {"id": True}}, {"unique_ids": {"id": "001"}},
    {"totals": {"amount": {"value": "30", "tolerance_typo": 100}}},
])
def test_malformed_contracts_never_pass(fixtures, contract):
    with pytest.raises(SkillError) as caught:
        validate_result_frame(read_table(fixtures / "sales.csv"), contract)
    assert caught.value.code == "SCHEMA_MISMATCH"


def test_explicit_valid_subset_with_exceptions(cli, tmp_path):
    source = tmp_path / "input.csv"
    source.write_text("id,amount\n001,10\n002,bad\n", encoding="utf-8")
    schema = tmp_path / "schema.json"
    schema.write_text(json.dumps({"required_columns": {"id": "string", "amount": "number"}}), encoding="utf-8")
    contract = tmp_path / "contract.json"
    contract.write_text(json.dumps({"row_count": 1, "totals": {"amount": 10}, "exception_count": 1}), encoding="utf-8")
    output = tmp_path / "subset.xlsx"
    report = cli("export_report", source, "--schema", schema, "--contract", contract,
                 "--output", output, "--allow-exceptions")
    assert report["validation_status"] == "passed_with_exceptions"
    assert report["invalid_rows"] == 1
    assert read_table(output, sheet="Exceptions").original_value.tolist() == ["bad"]
    cli("validate_result", output, "--sheet", "Cleaned", "--contract", contract,
        "--exceptions-file", output, "--exceptions-sheet", "Exceptions")


def test_full_xlsx_input_workflow(cli, fixtures, tmp_path):
    input_path = tmp_path / "input.xlsx"
    write_tables({"Sales": read_table(fixtures / "sales.csv")}, input_path)
    output = tmp_path / "processed.xlsx"
    report = cli("export_report", input_path, "--sheet", "Sales", "--schema", ROOT / "specs" / "schema.example.json",
                 "--contract", fixtures / "result.json", "--output", output)
    assert report["valid_rows"] == 2
    cleaned = read_table(output, sheet="Cleaned")
    assert cleaned.employee_id.tolist() == ["001234", "001235"]
    assert cleaned.source_sheet.tolist() == ["Sales", "Sales"]
