"""Conservative CSV/XLSX input, source lineage, and atomic exports."""
from __future__ import annotations

import csv
import os
import tempfile
from pathlib import Path
from typing import Any, Iterator

import pandas as pd
from openpyxl import load_workbook

from .errors import SkillError

PROVENANCE = ["source_file", "source_sheet", "source_row"]
FORMATS = {".csv", ".xlsx"}


def check_input(path: str | Path) -> Path:
    path = Path(path).resolve()
    if not path.is_file():
        raise SkillError("FILE_NOT_FOUND", "Input file not found.", file=str(path))
    if path.suffix.lower() not in FORMATS:
        raise SkillError("UNSUPPORTED_FORMAT", "Supported formats: CSV, XLSX.", file=str(path))
    return path


def check_headers(headers: list[Any], path: Path) -> None:
    if not headers or any(h is None or not str(h).strip() for h in headers):
        raise SkillError("HEADER_NOT_FOUND", "Headers must be nonempty.", file=str(path))
    if len(set(map(str, headers))) != len(headers):
        raise SkillError("SCHEMA_MISMATCH", "Duplicate column headers are ambiguous.", file=str(path))


def workbook_sheets(path: str | Path) -> list[str]:
    path = check_input(path)
    if path.suffix.lower() != ".xlsx":
        raise SkillError("UNSUPPORTED_FORMAT", "Expected XLSX workbook.", file=str(path))
    try:
        with pd.ExcelFile(path, engine="openpyxl") as book:
            return book.sheet_names
    except (ValueError, OSError, EOFError) as exc:
        raise SkillError("EXCEL_READ_ERROR", str(exc), file=str(path)) from exc


def prepare_input(path: Path, sheet: str | None, encoding: str, sep: str,
                  header: int) -> tuple[str | None, list[str]]:
    """Reject ambiguous headers, malformed CSV rows, and formula-bearing XLSX."""
    if header < 0 or len(sep) != 1:
        raise SkillError("INVALID_ARGUMENT", "header must be >= 0; delimiter must be one character.")
    if path.suffix.lower() == ".csv":
        if sheet is not None:
            raise SkillError("INVALID_ARGUMENT", "--sheet applies only to XLSX.")
        with path.open(encoding=encoding, newline="") as stream:
            reader = csv.reader(stream, delimiter=sep, strict=True)
            for _ in range(header):
                next(reader, None)
            headers = next(reader, [])
            check_headers(headers, path)
            for ordinal, record in enumerate(reader, start=header + 2):
                if len(record) != len(headers):
                    raise SkillError("CSV_PARSE_ERROR", "Record width differs from header.",
                                     file=str(path), row=ordinal)
        return None, headers
    with path.open("rb") as stream:
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            selected = sheet if sheet is not None else book.sheetnames[0]
            if selected not in book.sheetnames:
                raise SkillError("SHEET_NOT_FOUND", "Sheet does not exist.", file=str(path), sheet=selected)
            rows = book[selected].iter_rows(min_row=header + 1)
            cells = next(rows, ())
            headers = [c.value for c in cells]
            check_headers(headers, path)
            for cells in rows:
                for cell in cells:
                    if cell.data_type == "f":
                        raise SkillError("EXCEL_READ_ERROR", "Formula cells require an explicitly prepared values-only input.",
                                         file=str(path), sheet=selected, cell=cell.coordinate)
            return selected, list(map(str, headers))
        finally:
            book.close()


def iter_tables(path: str | Path, *, sheet: str | None = None,
                encoding: str = "utf-8-sig", sep: str = ",", header: int = 0,
                usecols: list[str] | None = None, chunksize: int = 100_000) -> Iterator[pd.DataFrame]:
    """Read raw fields as nullable strings; only empty fields count as missing."""
    path = check_input(path)
    if chunksize < 1:
        raise SkillError("INVALID_ARGUMENT", "chunksize must be positive.")
    try:
        selected, headers = prepare_input(path, sheet, encoding, sep, header)
        if usecols and not set(usecols).issubset(headers):
            raise SkillError("SCHEMA_MISMATCH", "Requested columns not found.", file=str(path))
        kwargs = dict(dtype="string", keep_default_na=False, na_values=[""],
                      header=header, usecols=usecols)
        if path.suffix.lower() == ".csv":
            with pd.read_csv(path, encoding=encoding, sep=sep, skip_blank_lines=False,
                             on_bad_lines="error", chunksize=chunksize, **kwargs) as reader:
                for frame in reader:
                    frame.attrs.update(source_file=str(path), source_sheet=None, header=header)
                    yield frame
        else:
            frame = pd.read_excel(path, sheet_name=selected, engine="openpyxl", **kwargs)
            frame.columns = frame.columns.map(str)
            frame.attrs.update(source_file=str(path), source_sheet=selected, header=header)
            yield frame
    except SkillError:
        raise
    except UnicodeError as exc:
        raise SkillError("ENCODING_ERROR", str(exc), file=str(path)) from exc
    except (csv.Error, pd.errors.ParserError, pd.errors.EmptyDataError) as exc:
        raise SkillError("CSV_PARSE_ERROR", str(exc), file=str(path)) from exc
    except (ValueError, OSError, EOFError) as exc:
        code = "CSV_PARSE_ERROR" if path.suffix.lower() == ".csv" else "EXCEL_READ_ERROR"
        raise SkillError(code, str(exc), file=str(path)) from exc


def read_table(path: str | Path, **kwargs: Any) -> pd.DataFrame:
    frames = list(iter_tables(path, **kwargs))
    if not frames:
        raise SkillError("HEADER_NOT_FOUND", "No readable table.", file=str(path))
    attrs = frames[0].attrs.copy()
    frame = pd.concat(frames, ignore_index=True)
    frame.attrs = attrs
    return frame


def add_provenance(frame: pd.DataFrame) -> pd.DataFrame:
    """Keep existing complete lineage; reject partially populated reserved fields."""
    present = set(PROVENANCE) & set(frame.columns)
    if present:
        if present != set(PROVENANCE):
            raise SkillError("SCHEMA_MISMATCH", "Partial provenance columns; rename reserved business fields.")
        row = pd.to_numeric(frame["source_row"], errors="coerce")
        if frame["source_file"].isna().any() or row.isna().any() or (row < 1).any() or (row % 1 != 0).any():
            raise SkillError("SCHEMA_MISMATCH", "Invalid existing provenance.")
        return frame.copy()
    result = frame.copy()
    result["source_file"] = frame.attrs.get("source_file", "in-memory")
    result["source_sheet"] = frame.attrs.get("source_sheet", pd.NA)
    result["source_row"] = range(frame.attrs.get("header", 0) + 2,
                                  frame.attrs.get("header", 0) + 2 + len(frame))
    return result


def write_tables(tables: dict[str, pd.DataFrame], output: str | Path, *,
                 sources: list[str | Path] | None = None, overwrite: bool = False,
                 spreadsheet_safe_csv: bool = False) -> dict[str, Any]:
    """Publish a fully written temporary file; never overwrite a source."""
    output = Path(output).resolve()
    if output.suffix.lower() not in FORMATS:
        raise SkillError("UNSUPPORTED_FORMAT", "Output must be CSV or XLSX.", file=str(output))
    for source in sources or []:
        source = Path(source).resolve()
        if source == output or (output.exists() and source.exists() and os.path.samefile(source, output)):
            raise SkillError("EXPORT_ERROR", "Source overwrite is prohibited; choose a new output path.", file=str(output))
    if output.exists() and not overwrite:
        raise SkillError("EXPORT_ERROR", "Output exists; use --overwrite only for an intended output replacement.", file=str(output))
    if not tables or (output.suffix.lower() == ".csv" and len(tables) != 1):
        raise SkillError("EXPORT_ERROR", "CSV requires exactly one table.")
    names = list(tables)
    if output.suffix.lower() == ".xlsx":
        if len({name.lower() for name in names}) != len(names):
            raise SkillError("EXPORT_ERROR", "Sheet names must be unique ignoring case.")
        for name, frame in tables.items():
            if not name or len(name) > 31 or any(c in name for c in "[]:*?/\\") or name.startswith("'") or name.endswith("'"):
                raise SkillError("EXPORT_ERROR", "Invalid sheet name.", sheet=name)
            if len(frame) > 1_048_575 or len(frame.columns) > 16_384:
                raise SkillError("EXPORT_ERROR", "Table exceeds XLSX dimensions.", sheet=name)
            if any(isinstance(v, str) and len(v) > 32_767 for v in frame.to_numpy().flat):
                raise SkillError("EXPORT_ERROR", "Cell text exceeds Excel length limit.", sheet=name)
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=".pandas-", suffix=output.suffix, dir=output.parent)
    os.close(fd)
    temp = Path(temporary)
    escaped = 0
    try:
        if output.suffix.lower() == ".csv":
            frame = next(iter(tables.values())).copy()
            if spreadsheet_safe_csv:
                def escape(value: Any) -> Any:
                    nonlocal escaped
                    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                        escaped += 1
                        return "'" + value
                    return value
                frame = frame.map(escape)
                frame.columns = [escape(str(c)) for c in frame.columns]
            frame.to_csv(temp, index=False, encoding="utf-8-sig")
        else:
            with pd.ExcelWriter(temp, engine="xlsxwriter", engine_kwargs={"options": {
                "strings_to_formulas": False, "strings_to_urls": False, "strings_to_numbers": False,
            }}) as writer:
                for name, frame in tables.items():
                    frame.to_excel(writer, sheet_name=name, index=False)
        if overwrite:
            os.replace(temp, output)
        else:
            # Hard-link publication fails atomically if the output appeared meanwhile.
            os.link(temp, output)
        return {"output": str(output), "rows": {name: len(frame) for name, frame in tables.items()},
                "csv_formula_cells_escaped": escaped}
    except (OSError, ValueError, TypeError) as exc:
        raise SkillError("EXPORT_ERROR", str(exc), file=str(output)) from exc
    finally:
        temp.unlink(missing_ok=True)
